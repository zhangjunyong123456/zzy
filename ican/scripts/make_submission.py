# -*- coding: utf-8 -*-
"""生成 iCAN 参赛提交包：截图副本 +「截图+文字说明」Word 文档。一次性脚本，运行后可删除。"""
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "screenshots"
OUT = ROOT / "submission"
OUT.mkdir(exist_ok=True)

SHOTS = [
    ("01-landing.png", "图1 产品落地页（首屏）",
     "品牌主张「一颗会成长的 AI，陪你走完大学的每一段路」，一句话点明五大场景（学习·科研·竞赛·求职·校园）一个入口；"
     "右下角为可交互桌面宠物「小桃」，点击走动、双击发起对话，体现产品的陪伴感与低门槛入口。"),
    ("01b-landing-features.png", "图2 落地页核心机制说明区",
     "以「路由 → 并行 → 汇总」三步讲解多 Agent 工作流——主 Agent 听懂用户意图并拆解任务，"
     "专业 Agent 同频并行生成，最终汇总为一份可执行的成长行动建议。"),
    ("02-login.png", "图3 登录/注册页",
     "支持账号登录与游客模式；登录后 AI 会结合个人画像给出更贴合的成长建议（密码 PBKDF2 加密存储，会话与消息数据按用户隔离）。"),
    ("03-chat-scenes.png", "图4 对话首页 · 场景选择",
     "登录后进入主界面，五大方向场景卡片一目了然；点击卡片即进入对应方向并获得引导提示，"
     "左侧为历史会话列表，底部建议问题一键发送，零打字即可开始体验。"),
    ("04-ai-multi-agent.png", "图5 AI 实际输出 · 多 Agent 并行协同（核心功能）",
     "输入「我要参加 iCAN 比赛，帮我全面规划」后，主 Agent 实际调用 DeepSeek 路由判定为 multi 任务，"
     "并行调度竞赛/科研/学习三个 Agent 同时流式作答（顶部链路图实时展示调度过程），"
     "最终生成金色「主 Agent 汇总卡——iCAN 综合成长行动建议」，含按优先级排序的下一步行动清单。"),
    ("05-model-config.png", "图6 个人中心 · 模型服务配置",
     "三供应商卡片（DeepSeek/硅基流动/智谱），支持在线保存 API Key、切换当前模型（当前使用 deepseek-flash），"
     "Key 仅存本机 backend/.env，保存即生效无需重启。"),
    ("06-docs.png", "图7 知识库 · PDF 即传即问",
     "拖拽上传 PDF（≤20MB/300 页），按场景归类（课程资料/科研论文/简历）；"
     "截图中 test_upload.pdf 已完成解析（状态「就绪」，3 页解析为 3 个分块并向量化入库），"
     "上传后即可在对话中针对该文档提问，回答带「（来源：第 N 页）」页码标注。"),
]

# 1) 复制截图副本
for name, _, _ in SHOTS:
    shutil.copy2(SRC / name, OUT / name)

# 2) 生成 Word 文档
doc = Document()
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)

title = doc.add_heading("UniGrow Agent（智学成长Agent）· 产品核心页面截图说明", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(
    "面向大学生全生命周期的多智能体成长平台：一个聊天入口，覆盖学习·科研·竞赛·求职·校园五大场景。"
    "以下截图均来自实际运行的系统，其中图 5 为 AI 实际输出界面（真实大模型流式生成，非示意图）。"
)

for name, cap, desc in SHOTS:
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(str(SRC / name), width=Cm(15.5))
    cp = doc.add_paragraph()
    r = cp.add_run(cap)
    r.bold = True
    dp = doc.add_paragraph(desc)
    doc.add_paragraph()

doc.save(str(OUT / "UniGrow-Agent-产品截图说明.docx"))
print(f"提交包已生成：{OUT}")
for f in sorted(OUT.iterdir()):
    print(f"  {f.name}  ({f.stat().st_size / 1024:.0f} KB)")
