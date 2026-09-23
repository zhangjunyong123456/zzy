"""附件服务：上传落盘、按类型提取文本（图片走视觉模型）、归属绑定与注入文本拼接。"""
import asyncio
import base64
import json
import logging
import uuid
from datetime import datetime, timezone
from pathlib import Path

from app.config import settings
from app.database import get_conn

logger = logging.getLogger(__name__)

MAX_TEXT = 8000  # 单个附件提取文本上限（注入时同样按此截断）

# 扩展名 → (mime, kind)
_FILE_TYPES = {
    ".png": ("image/png", "image"),
    ".jpg": ("image/jpeg", "image"),
    ".jpeg": ("image/jpeg", "image"),
    ".webp": ("image/webp", "image"),
    ".pdf": ("application/pdf", "file"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "file"),
    ".pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "file"),
    ".txt": ("text/plain", "file"),
    ".md": ("text/markdown", "file"),
}

ACCEPT = ",".join(k.lstrip(".") for k in _FILE_TYPES)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _owned(row, uid: str | None) -> bool:
    return (row["user_id"] or None) == (uid or None)


def att_dir() -> Path:
    d = Path(settings.upload_dir) / "attachments"
    d.mkdir(parents=True, exist_ok=True)
    return d


def resolve_type(filename: str) -> tuple[str, str] | None:
    """按扩展名返回 (mime, kind)；不认识返回 None。"""
    ext = "." + filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    return _FILE_TYPES.get(ext)


# ---------- 提取 ----------

def _clip(text: str) -> str:
    text = text.strip()
    if len(text) > MAX_TEXT:
        text = text[:MAX_TEXT] + "（已截断）"
    return text


def _extract_doc(path: Path, ext: str) -> str:
    """同步解析：pdf / docx / pptx / txt / md（阻塞调用须包 asyncio.to_thread）。"""
    try:
        if ext == ".pdf":
            from app.rag.parser import parse_pdf

            pages = parse_pdf(path.read_bytes())
            return _clip("\n\n".join(f"[第{p}页]\n{t}" for p, t in pages))
        if ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))
            return _clip("\n".join(parts))
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))

            def shape_texts(shapes):
                """递归展开组合形状，收集文本与表格。"""
                for shp in shapes:
                    if shp.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                        yield from shape_texts(shp.shapes)
                        continue
                    if getattr(shp, "has_table", False) and shp.has_table:
                        for row in shp.table.rows:
                            cells = [c.text.strip() for c in row.cells]
                            if any(cells):
                                yield "\t".join(cells)
                        continue
                    if getattr(shp, "has_text_frame", False):
                        t = shp.text_frame.text.strip()
                        if t:
                            yield t

            parts = []
            for i, slide in enumerate(prs.slides, 1):
                texts = list(shape_texts(slide.shapes))
                if texts:
                    parts.append(f"[第{i}页]\n" + "\n".join(texts))
            return _clip("\n\n".join(parts))
        # txt / md
        data = path.read_bytes()
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("gbk", errors="ignore")
        return _clip(text)
    except Exception as e:
        logger.warning("附件解析失败 %s: %s", path.name, e)
        return ""


async def extract_image(path: Path, mime: str) -> str:
    """视觉模型提取图片内容；无可用 Key（用户/全局）或调用失败返回空串。

    Key 选择逻辑在 get_vision_llm 内：用户智谱 Key → 全局 Key（仅非独占模式）→ None 跳过。
    """
    from app.agents.llm import get_vision_llm

    llm = get_vision_llm()
    if llm is None:
        return ""
    from langchain_core.messages import HumanMessage

    b64 = base64.b64encode(path.read_bytes()).decode()
    msg = HumanMessage(
        content=[
            {
                "type": "text",
                "text": "请完整提取并描述这张图片中的所有内容，包括文字、图表、公式、代码等，用中文输出，不要遗漏。",
            },
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
        ]
    )
    try:
        resp = await llm.ainvoke([msg])
        text = resp.content if isinstance(resp.content, str) else ""
    except Exception as e:
        logger.warning("视觉模型提取图片失败: %s", e)
        return ""
    return _clip(text)


async def extract_text(path: Path, ext: str, mime: str, kind: str) -> str:
    if kind == "image":
        return await extract_image(path, mime)
    return await asyncio.to_thread(_extract_doc, path, ext)


# ---------- CRUD / 绑定 ----------

def create_attachment(
    user_id: str | None, session_id: str | None, filename: str,
    mime: str, kind: str, path: str, size: int, extracted_text: str,
) -> dict:
    att_id = uuid.uuid4().hex
    with get_conn() as conn:
        conn.execute(
            "INSERT INTO attachments (id, user_id, session_id, message_id, filename, mime, "
            "kind, path, size, extracted_text, created_at) VALUES (?, ?, ?, NULL, ?, ?, ?, ?, ?, ?, ?)",
            (att_id, user_id, session_id, filename, mime, kind, path, size, extracted_text, _now()),
        )
    return get_attachment(att_id)


def get_attachment(att_id: str) -> dict | None:
    with get_conn() as conn:
        row = conn.execute("SELECT * FROM attachments WHERE id = ?", (att_id,)).fetchone()
    return dict(row) if row else None


def bind_attachments(
    message_id: str, session_id: str, att_ids: list[str], uid: str | None
) -> list[dict]:
    """把本轮上传的附件绑定到用户消息；越权/不属于本会话的 id 静默忽略。返回绑定清单。"""
    bound: list[dict] = []
    for att_id in att_ids:
        att = get_attachment(att_id)
        if not att or not _owned(att, uid):
            continue
        if att["session_id"] not in (None, session_id):
            continue
        with get_conn() as conn:
            conn.execute(
                "UPDATE attachments SET session_id = ?, message_id = ? WHERE id = ?",
                (session_id, message_id, att_id),
            )
        bound.append({"id": att_id, "filename": att["filename"], "kind": att["kind"], "mime": att["mime"]})
    if bound:
        with get_conn() as conn:
            conn.execute(
                "UPDATE messages SET attachments = ? WHERE id = ?",
                (json.dumps(bound, ensure_ascii=False), message_id),
            )
    return bound


def build_attachments_text(bound: list[dict]) -> str:
    """绑定清单 → 注入文本「[附件：名]\n提取内容」，总长 MAX_TEXT 截断。"""
    blocks: list[str] = []
    total = 0
    for item in bound:
        att = get_attachment(item["id"])
        if not att or not att["extracted_text"].strip():
            continue
        block = f"[附件：{att['filename']}]\n{att['extracted_text']}"
        if total + len(block) > MAX_TEXT:
            remain = MAX_TEXT - total
            if remain > 50:
                blocks.append(block[:remain] + "（已截断）")
            break
        blocks.append(block)
        total += len(block)
    return "\n\n".join(blocks)
